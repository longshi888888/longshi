"""Bounded office tools: selected input folder, new output files, no shell."""
from pathlib import Path
from decimal import Decimal,localcontext
from datetime import datetime
import ast
import csv
import fnmatch
import hashlib
import io
import json
import math
import os
import re
import threading
import uuid
import zipfile

SUPPORTED={'.txt','.md','.csv','.json','.docx','.xlsx','.pdf'}
SKIP={'.git','.svn','node_modules','__pycache__','.venv','venv','runtime'}
SENSITIVE={'credential.bin','credentials.json','secrets.json','config.toml','settings.json'}
MEMORY_KEYS={'称呼','报告风格','输出偏好','工作习惯'}

def demo_texts():
    return {
        '示例会议记录.md':'# 合成会议记录\n\n以下日期、人名、事项全部虚构，不是公司真实业务。\n\n会议日期：2099-01-05\n主题：演示项目周会\n\n1. 示例负责人甲：2099-01-08前完成资料清单。\n2. 示例负责人乙：2099-01-09前整理费用表。\n3. 待确认：下次会议的具体时间，记录中没有结论。\n\n费用参考见示例费用.csv；没有付款凭证。\n',
        '示例费用.csv':'项目,金额,说明\n演示材料,100.50,全部虚构\n演示运输,20.25,全部虚构\n演示服务,30.00,全部虚构\n',
        '使用说明.txt':'这里全部是虚构示例资料。可以让小隆列文件、搜索日期，或接模型后整理会议纪要和费用汇总。请选择自己的办公文件夹后再处理实际资料。'}

def prepare_demo(folder):
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
    for name,content in demo_texts().items():
        path=folder/name
        if not path.exists():path.write_text(content,encoding='utf-8')

def calculate_decimal(expression):
    if not isinstance(expression,str) or len(expression)>4000:raise ValueError('计算式过长')
    try:tree=ast.parse(expression,mode='eval')
    except SyntaxError:raise ValueError('计算式不合法') from None
    if len(list(ast.walk(tree)))>1200:raise ValueError('计算式过于复杂')
    def visit(node):
        if isinstance(node,ast.Constant) and type(node.value) in (int,float):
            value=Decimal(str(node.value))
            if not value.is_finite() or value.adjusted()>30:raise ValueError('数字超出范围')
            return value
        if isinstance(node,ast.List) and len(node.elts)<=500:return [visit(x) for x in node.elts]
        if isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.UAdd,ast.USub)):
            value=visit(node.operand);return value if isinstance(node.op,ast.UAdd) else -value
        if isinstance(node,ast.BinOp):
            left,right=visit(node.left),visit(node.right)
            if not isinstance(left,Decimal) or not isinstance(right,Decimal):raise ValueError('算术只接受数字')
            if isinstance(node.op,ast.Add):return left+right
            if isinstance(node.op,ast.Sub):return left-right
            if isinstance(node.op,ast.Mult):return left*right
            if isinstance(node.op,ast.Div):return left/right
        if isinstance(node,ast.Call) and isinstance(node.func,ast.Name) and node.func.id in ('sum','min','max','len') and len(node.args)==1 and not node.keywords:
            values=visit(node.args[0])
            if not isinstance(values,list) or not all(isinstance(v,Decimal) for v in values):raise ValueError('统计函数只接受数字列表')
            if node.func.id=='sum':return sum(values,Decimal('0'))
            if node.func.id=='len':return Decimal(len(values))
            if not values:raise ValueError('不能计算空列表的最值')
            return min(values) if node.func.id=='min' else max(values)
        raise ValueError('只支持数字、四则运算及sum/min/max/len，不执行代码')
    try:
        with localcontext() as ctx:
            ctx.prec=28;value=visit(tree.body)
        if not isinstance(value,Decimal) or not value.is_finite() or value.adjusted()>30:raise ValueError('计算结果超出范围')
        return format(value,'f')
    except ArithmeticError:raise ValueError('计算失败，请检查除零或数字范围') from None

class OfficeTools:
    def __init__(self,store,stage,stop=None):
        self.store=store;self.root=Path(store.config['workspace']).resolve()
        self.output=store.folder/'outputs';self.output.mkdir(parents=True,exist_ok=True)
        self.stage=stage;self.stop=stop or threading.Event();self.sources={};self.artifacts=[];self.trace=[];self.lock=threading.RLock()
    def check(self):
        if self.stop.is_set():raise InterruptedError('任务已停止')
    def begin(self,label):
        self.check();step={'label':label,'state':'running'};self.trace.append(step);self.stage('working',label,{'steps':list(self.trace),'artifacts':list(self.artifacts),'sources':list(self.sources.values())});return step
    def finish(self,step):
        step['state']='done';self.stage('working',step['label'],{'steps':list(self.trace),'artifacts':list(self.artifacts),'sources':list(self.sources.values())})
    def path(self,relative):
        if not isinstance(relative,str) or len(relative)>500 or '\x00' in relative:raise ValueError('文件路径不合法')
        path=Path(relative or '.')
        if path.is_absolute() or path.drive or '..' in path.parts:raise ValueError('只允许所选工作文件夹内的相对路径')
        resolved=(self.root/path).resolve()
        if not resolved.is_relative_to(self.root):raise ValueError('文件链接指向工作范围以外')
        if any(part.startswith('.') or part in SKIP for part in path.parts if part!='.') or resolved.name.lower() in SENSITIVE:raise ValueError('该文件不作为办公资料读取')
        return resolved
    def _reference(self,path,location):
        relative=path.relative_to(self.root).as_posix();digest=hashlib.sha256(path.read_bytes()).hexdigest()
        ref={'file':relative,'location':location,'sha256':digest,'label':'用户选择的工作资料','workspace_id':self.store.config['workspace_id']}
        self.sources[relative]=ref;return ref
    def _text(self,path):
        raw=path.read_bytes()
        if raw[:2] in (b'\xff\xfe',b'\xfe\xff'):return raw.decode('utf-16')
        try:return raw.decode('utf-8-sig')
        except UnicodeDecodeError:return raw.decode('gb18030')
    def _check_zip(self,path):
        with zipfile.ZipFile(path) as archive:
            infos=archive.infolist()
            if len(infos)>10000 or sum(i.file_size for i in infos)>50_000_000:raise ValueError('文档展开后过大，本次未读取')
    def list_files(self,folder:str='',pattern:str='*',recursive:bool=True)->dict:
        """列出用户所选工作文件夹内的资料。返回相对路径、类型和大小；不扫描整台电脑。"""
        step=self.begin('查看工作资料清单');base=self.path(folder)
        if not base.is_dir():raise ValueError('没有这个子文件夹')
        if len(pattern)>100:raise ValueError('文件筛选条件过长')
        result=[];scanned=0;truncated=False
        def paths():
            if not recursive:yield from base.iterdir();return
            for folder,dirs,names in os.walk(base,followlinks=False):
                approved=[]
                for directory in dirs:
                    try:self.path(str((Path(folder)/directory).relative_to(self.root)))
                    except (ValueError,OSError):continue
                    approved.append(directory)
                dirs[:]=approved
                for name in names:yield Path(folder)/name
        for path in paths():
            self.check();scanned+=1
            if scanned>3000 or len(result)>=300:truncated=True;break
            relative=path.relative_to(self.root)
            try:allowed=self.path(str(relative))
            except (ValueError,OSError):continue
            if allowed.is_file() and allowed.suffix.lower() in SUPPORTED and fnmatch.fnmatch(path.name,pattern):
                result.append({'path':relative.as_posix(),'type':allowed.suffix.lower(),'bytes':allowed.stat().st_size})
        self.finish(step);return {'files':result,'truncated':truncated,'scope':'仅用户选定文件夹'}
    def read_document(self,path:str,start:int=1,limit:int=120,sheet:str='')->dict:
        """读取TXT/MD/CSV/JSON、DOCX、XLSX或可选文字PDF，返回定位与哈希。start/limit为行、表格行或PDF页；不会执行Excel公式、宏或文档指令。"""
        step=self.begin('读取资料：'+str(path)[:80]);file=self.path(path)
        if not file.is_file() or file.suffix.lower() not in SUPPORTED:raise ValueError('仅支持文本、DOCX、XLSX和可提取文字PDF')
        if file.stat().st_size>12_000_000:raise ValueError('文件超过12MB，本次未读取')
        if start<1 or limit<1:raise ValueError('读取起点与长度必须为正数')
        before=hashlib.sha256(file.read_bytes()).hexdigest()
        limit=min(limit,200);suffix=file.suffix.lower();self.check()
        if suffix=='.xlsx':
            from openpyxl import load_workbook
            self._check_zip(file);book=load_workbook(file,read_only=True,data_only=False)
            try:
                if sheet and sheet not in book.sheetnames:raise ValueError('没有这个工作表，可用：'+', '.join(book.sheetnames))
                ws=book[sheet] if sheet else book.worksheets[0];rows=[]
                from openpyxl.utils import get_column_letter
                for ri,row in enumerate(ws.iter_rows(min_row=start,max_row=min(start+limit-1,ws.max_row),max_col=min(ws.max_column,40)),start):
                    self.check();rows.append([{'cell':f'{get_column_letter(ci)}{ri}','value':str(c.value)[:2000] if c.value is not None else '', 'formula':getattr(c,'data_type','')=='f'} for ci,c in enumerate(row,1)])
                ref=self._reference(file,f'{ws.title}!行{start}—{start+len(rows)-1}')
                result={'rows':rows,'sheets':book.sheetnames,'source':ref,'note':'公式为原文，未重算；仅所选行列'}
            finally:book.close()
        else:
            if suffix=='.pdf':
                from pypdf import PdfReader
                reader=PdfReader(str(file))
                if reader.is_encrypted:raise ValueError('PDF已加密，未读取')
                count=min(limit,5);lines=[]
                for page in range(start-1,min(start-1+count,len(reader.pages))):
                    self.check();lines.append(f'[第{page+1}页]\n'+(reader.pages[page].extract_text() or ''))
                text='\n'.join(lines)
                if len(re.sub(r'\[第\d+页\]','',text).strip())<3:raise ValueError('PDF没有可提取文字，可能需要另做OCR')
                location=f'第{start}—{min(start-1+count,len(reader.pages))}页'
            else:
                if suffix=='.docx':
                    from docx import Document
                    self._check_zip(file);doc=Document(file);lines=[p.text for p in doc.paragraphs]
                    for i,table in enumerate(doc.tables,1):
                        lines += [f'[表{i}行{j}] '+' | '.join(cell.text for cell in row.cells) for j,row in enumerate(table.rows,1)]
                else:lines=self._text(file).splitlines()
                selected=lines[start-1:start-1+limit];text='\n'.join(f'{start+i}: {line}' for i,line in enumerate(selected));location=f'行/段落{start}—{start+len(selected)-1}'
            result={'text':text[:24000],'truncated':len(text)>24000,'source':self._reference(file,location)}
        if result['source']['sha256']!=before:raise ValueError('资料读取期间发生变化，请重新读取')
        self.finish(step);return result
    def search_documents(self,query:str,max_results:int=8)->dict:
        """在选定工作资料中查找关键词，返回实际片段和出处。最多扫描80个支持的文件；不是互联网搜索。"""
        if not query.strip() or len(query)>200:raise ValueError('请输入1至200字关键词')
        if not 1<=max_results<=15:raise ValueError('最多返回1至15个匹配片段')
        step=self.begin('查找资料：'+query[:50]);found=[];errors=[]
        entries=self.list_files()['files'][:80]
        for entry in entries:
            self.check()
            try:
                document=self.read_document(entry['path'],limit=200)
                text=document.get('text') or json.dumps(document.get('rows',[]),ensure_ascii=False)
                for match in re.finditer(re.escape(query),text,re.I):
                    found.append({'text':text[max(0,match.start()-100):match.end()+180],'source':document['source']})
                    if len(found)>=min(max_results,15):break
            except (ValueError,OSError,zipfile.BadZipFile) as e:errors.append({'file':entry['path'],'reason':str(e)[:150]})
            if len(found)>=min(max_results,15):break
        self.finish(step);return {'matches':found,'unreadable':errors,'scope':'每文件有限行/页；没有命中不代表全文件不存在'}
    def calculate(self,expression:str)->dict:
        """使用Decimal计算四则运算或sum/min/max/len数字列表，避免靠语言模型心算。不运行Python或任意代码。"""
        step=self.begin('核对数字');result=calculate_decimal(expression);self.finish(step);return {'expression':expression,'result':result,'method':'Decimal，28位精度'}
    def artifact_info(self,artifact_id):
        if not isinstance(artifact_id,str) or not re.fullmatch('[a-f0-9]{32}',artifact_id):raise ValueError('成果编号不合法')
        meta=json.loads((self.output/(artifact_id+'.json')).read_text(encoding='utf-8'))
        if not isinstance(meta,dict) or not isinstance(meta.get('name'),str):raise ValueError('成果清单不完整')
        file=(self.output/meta['name']).resolve()
        if not file.is_relative_to(self.output.resolve()) or file.suffix not in ('.docx','.xlsx','.md','.txt','.csv','.json'):raise ValueError('成果路径不合法')
        if hashlib.sha256(file.read_bytes()).hexdigest()!=meta.get('sha256'):raise ValueError('成果已经变化，请重新核对')
        return file,meta
    def list_artifacts(self)->dict:
        """列出小隆最近实际生成的20个成果及编号，便于继续修改上一份文档。不是工作文件夹的原件清单。"""
        step=self.begin('查找已有成果');items=[]
        for file in sorted(self.output.glob('*.json'),key=lambda p:p.stat().st_mtime,reverse=True)[:20]:
            try:items.append(self.artifact_info(file.stem)[1])
            except (OSError,ValueError,KeyError):continue
        self.finish(step);return {'artifacts':items}
    def read_artifact(self,artifact_id:str,start:int=1,limit:int=120,sheet:str='')->dict:
        """读取小隆已生成且哈希一致的成果，供后续修改。先list_artifacts取编号；修改时须调用生成工具另存新版本。"""
        step=self.begin('读取已有成果');file,meta=self.artifact_info(artifact_id)
        reader=OfficeTools(self.store,lambda *args:None,self.stop);reader.root=self.output.resolve()
        result=reader.read_document(file.name,start,limit,sheet)
        reference=dict(result['source']);reference.pop('workspace_id',None)
        reference.update(artifact_id=artifact_id,label='小隆已生成的成果')
        result['source']=reference;self.sources['artifact:'+artifact_id]=reference
        self.finish(step);return result
    def _save(self,name,suffix,raw):
        self.check()
        if hasattr(self,'before_write'):self.before_write('保存成果：'+str(name)[:80])
        if not isinstance(name,str) or not name.strip() or len(name)>100 or any(c in name for c in '/\\:\x00\r\n<>|?*'):raise ValueError('成果名称只能是普通文件名，不能带路径')
        stem=Path(name).stem.strip(' .')
        if not stem or stem.upper() in {'CON','PRN','AUX','NUL',*(f'COM{i}' for i in range(1,10)),*(f'LPT{i}' for i in range(1,10))}:raise ValueError('成果文件名不合法')
        if len(raw)>10_000_000:raise ValueError('成果文件过大')
        aid=uuid.uuid4().hex;filename=stem+'-'+aid[:8]+suffix;file=self.output/filename
        with file.open('xb') as f:f.write(raw)
        artifact={'id':aid,'name':filename,'bytes':len(raw),'sha256':hashlib.sha256(raw).hexdigest(),'created_at':datetime.now().astimezone().isoformat(timespec='seconds'),'session_id':getattr(self.store,'session_id','legacy')}
        # Only this directory's new files can be opened/downloaded through the API.
        (self.output/(aid+'.json')).write_text(json.dumps(artifact,ensure_ascii=False),encoding='utf-8')
        self.artifacts.append(artifact);return artifact
    def create_document(self,name:str,title:str,content:str,format:str='docx')->dict:
        """生成新的Word或Markdown办公草稿。content为正文，可用#标题分节；只写独立成果目录，不覆盖原件。只有工具成功返回才算生成。"""
        if format not in ('docx','md') or len(content)>100000 or len(title)>300:raise ValueError('正文过长或输出格式不支持')
        step=self.begin('生成文档：'+name[:60])
        if format=='md':raw=('# '+title+'\n\n'+content).encode('utf-8')
        else:
            from docx import Document
            from docx.shared import Pt
            doc=Document();doc.styles['Normal'].font.name='Microsoft YaHei';doc.styles['Normal'].font.size=Pt(11)
            doc.add_heading(title,0)
            for line in content.splitlines():
                self.check()
                match=re.match(r'^(#{1,3})\s+(.+)$',line)
                if match:doc.add_heading(match.group(2),min(len(match.group(1)),3))
                else:doc.add_paragraph(line)
            buffer=io.BytesIO();doc.save(buffer);raw=buffer.getvalue()
        artifact=self._save(name,'.'+format,raw);self.finish(step);return artifact
    def create_workbook(self,name:str,sheets:list[dict])->dict:
        """生成新Excel文件。sheets每项含title、columns数组、rows二维数组。数字请先用calculate核算；字符串永远不作为公式执行，不修改原工作簿。"""
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        if not sheets or len(sheets)>8:raise ValueError('请提供1至8张表')
        step=self.begin('生成表格：'+name[:60]);book=Workbook();book.remove(book.active)
        for spec in sheets:
            self.check();columns=spec.get('columns',[]);rows=spec.get('rows',[]);title=spec.get('title','工作表')
            if not isinstance(columns,list) or not 1<=len(columns)<=40 or not isinstance(rows,list) or len(rows)>5000:raise ValueError('表格行列超出范围')
            ws=book.create_sheet(re.sub(r'[\\/*?:\[\]]','_',str(title))[:31] or '工作表')
            def literal(value):
                if value is None:return ''
                if isinstance(value,(int,float)) and not isinstance(value,bool):
                    if not math.isfinite(value):raise ValueError('不能写入非有限数字')
                    return value
                text=str(value)
                if len(text)>10000:raise ValueError('单元格内容过长')
                return "'"+text if text.startswith(('=','+','-','@')) else text
            ws.append([literal(v) for v in columns])
            for row in rows:
                self.check()
                if not isinstance(row,list) or len(row)>len(columns):raise ValueError('表格行不匹配')
                ws.append([literal(v) for v in row])
            for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='657F51')
            for row in ws:
                for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
            for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=min(38,max(14,max(len(str(c.value or '')) for c in column[:100])+3))
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        buffer=io.BytesIO();book.save(buffer);artifact=self._save(name,'.xlsx',buffer.getvalue());self.finish(step);return artifact
    def remember_preference(self,key:str,value:str)->dict:
        """只保存用户明确要求记住的称呼、报告风格、输出偏好或工作习惯，不保存凭据或正式业务事实。用户可在界面查看和删除。"""
        if key not in MEMORY_KEYS or not value.strip() or len(value)>500:raise ValueError('仅能记住简短办公偏好')
        if re.search(r'(?i)(api.?key|bearer\s|sk-[a-z0-9]{12}|密码|口令|密钥)',value):raise ValueError('凭据请使用设置，不保存到偏好')
        if hasattr(self,'before_write'):self.before_write('保存明确要求的办公偏好')
        step=self.begin('记住办公偏好');self.store.remember(key,value.strip());self.finish(step);return {'saved':True,'key':key,'value':value.strip()}
    def recall_preferences(self)->dict:
        """查看用户已保存的有限办公偏好，不作为公司业务事实或制度。"""
        self.check();return dict(self.store.preferences)
