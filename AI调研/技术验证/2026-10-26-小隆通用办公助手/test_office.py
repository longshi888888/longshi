import hashlib
import io
import json
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
import urllib.request
import urllib.error
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer
from office_tools import OfficeTools,calculate_decimal
from office_agent import OfficeAgent
from storage import Store,RUNTIME
from server import Server

class ScriptedModel(BaseHTTPRequestHandler):
    def log_message(self,*args):pass
    def do_POST(self):
        self.server.payloads.append(json.loads(self.rfile.read(int(self.headers['Content-Length']))))
        time.sleep(self.server.delay)
        index=self.server.count;self.server.count+=1
        if index<len(self.server.script):
            name,args=self.server.script[index]
            message={'role':'assistant','content':None,'tool_calls':[{'id':f'call-{index}','type':'function','function':{'name':name,'arguments':json.dumps(args,ensure_ascii=False)}}]};finish='tool_calls'
        else:message={'role':'assistant','content':'已完成合成测试任务，实际成果见文件卡。'};finish='stop'
        if self.server.payloads[-1].get('stream'):
            try:
                self.send_response(200);self.send_header('Content-Type','text/event-stream');self.send_header('Connection','close');self.end_headers()
                def emit(delta,reason=None):
                    event={'id':'synthetic','object':'chat.completion.chunk','created':1,'model':'synthetic-office-stub','choices':[{'index':0,'delta':delta,'finish_reason':reason}]}
                    self.wfile.write(('data: '+json.dumps(event)+'\n\n').encode());self.wfile.flush()
                    time.sleep(getattr(self.server,'chunk_delay',0))
                if message.get('tool_calls'):
                    call=message['tool_calls'][0];function=call['function']
                    emit({'role':'assistant','tool_calls':[{'index':0,'id':call['id'],'type':'function','function':{'name':function['name'],'arguments':''}}]})
                    arguments=function['arguments']
                    for offset in range(0,len(arguments),12):emit({'tool_calls':[{'index':0,'function':{'arguments':arguments[offset:offset+12]}}]})
                else:
                    emit({'role':'assistant','reasoning_content':'SYNTHETIC-PRIVATE-THINKING'})
                    for offset in range(0,len(message['content']),4):emit({'content':message['content'][offset:offset+4]})
                emit({},finish)
                usage={'id':'synthetic','object':'chat.completion.chunk','created':1,'model':'synthetic-office-stub','choices':[],'usage':{'prompt_tokens':100,'completion_tokens':30,'total_tokens':130}}
                self.wfile.write(('data: '+json.dumps(usage)+'\n\ndata: [DONE]\n\n').encode());self.wfile.flush()
            except (BrokenPipeError,ConnectionResetError,ConnectionAbortedError):pass
            return
        raw=json.dumps({'id':'synthetic','object':'chat.completion','created':1,'model':'synthetic-office-stub','choices':[{'index':0,'finish_reason':finish,'message':message}],'usage':{'prompt_tokens':100,'completion_tokens':30,'total_tokens':130}}).encode()
        try:self.send_response(200);self.send_header('Content-Type','application/json');self.send_header('Content-Length',str(len(raw)));self.end_headers();self.wfile.write(raw)
        except (BrokenPipeError,ConnectionResetError):pass

class OfficeTests(unittest.TestCase):
    def setUp(self):
        RUNTIME.mkdir(parents=True,exist_ok=True);self.temp=tempfile.TemporaryDirectory(dir=RUNTIME,prefix='office-test-')
        self.store=Store(self.temp.name);self.tools=OfficeTools(self.store,lambda *args:None)
    def tearDown(self):self.temp.cleanup()
    def test_real_stream_reaches_http_before_completion_and_never_exposes_thinking(self):
        model=ThreadingHTTPServer(('127.0.0.1',0),ScriptedModel);model.count=0;model.payloads=[];model.delay=0;model.chunk_delay=.035;model.script=[]
        mt=threading.Thread(target=model.serve_forever,daemon=True);mt.start()
        self.store.save({'base_url':f'http://127.0.0.1:{model.server_port}','model':'synthetic-stream','api_key':'SYNTHETIC-STREAM','use_model':True})
        server=Server(self.store);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        request=urllib.request.Request(server.base+'ask',data=json.dumps({'question':'合成流式测试'}).encode(),headers={'Content-Type':'application/json'})
        try:
            jid=json.load(urllib.request.urlopen(request))['id'];updates=[]
            with urllib.request.urlopen(server.base+'events/'+jid,timeout=10) as response:
                self.assertIn('text/event-stream',response.headers['Content-Type'])
                for line in response:
                    if line.startswith(b'data: '):
                        job=json.loads(line[6:]);updates.append(job)
                        if job['state'] in ('done','error','cancelled'):break
            streaming=[j for j in updates if j['state']=='streaming' and j.get('stream_text')]
            self.assertGreaterEqual(len(streaming),2)
            self.assertLess(len(streaming[0]['stream_text']),len(updates[-1]['answer']['body']))
            self.assertEqual(updates[-1]['state'],'done');self.assertEqual(updates[-1]['answer']['body'],streaming[-1]['stream_text'])
            self.assertNotIn('PRIVATE-THINKING',json.dumps(updates));self.assertTrue(model.payloads[0]['stream'])
            self.assertEqual(len(self.store.history),1)
        finally:server.shutdown();server.server_close();thread.join(3);model.shutdown();model.server_close();mt.join(3)
    def test_unified_capabilities_use_actual_switch_and_do_not_intercept_tasks(self):
        agent=OfficeAgent();self.store.config['use_model']=True
        result=agent.run('你可以做什么？','office',self.store.config,'SYNTHETIC',[],lambda *args:None,self.store)
        self.assertEqual(result['kind'],'capabilities');self.assertIn('当前模型已启用',result['body'])
        self.assertNotIn('切换到',result['body']);self.assertNotIn('还没有连接',result['body'])
        self.store.config['use_model']=False
        result=agent.run('明明连上了为什么说没连上？','office',self.store.config,'SYNTHETIC',[],lambda *args:None,self.store)
        self.assertIn('当前模型未启用',result['body'])
        from business_agent import capability_answer
        self.assertIsNone(capability_answer('你可以做一份会议纪要吗',self.store.config,''))
    def test_legacy_chat_does_not_reuse_unconfigured_local_fallback(self):
        from business_agent import Agent
        history=[{'mode':'chat','question':'你好','answer':{'engine':'本地提示（未调用模型）','body':'尚未连接模型'}}]
        with patch('business_agent.completion',return_value={'content':'你好，我在。'}) as call:
            Agent().run('今天辛苦了','chat',{'use_model':True,'nickname':'你'},'SYNTHETIC',history,lambda *a:None)
            messages=call.call_args.args[2]
            self.assertEqual(len(messages),2);self.assertIn('模型已启用',messages[0]['content'])
    def test_api_draft_normalizes_official_name_without_mutation(self):
        self.store.save({'api_key':'SYNTHETIC-LOCAL-ONLY','model':'old-saved-id'})
        before=(self.store.folder/'settings.json').read_bytes()
        credential=(self.store.folder/'credential.bin').read_bytes()
        config,key=self.store.model_settings({'model':'DeepSeek V4.1 Flash'})
        self.assertEqual(config['model'],'deepseek-flash');self.assertEqual(key,'SYNTHETIC-LOCAL-ONLY')
        self.assertFalse(config['use_model']);self.assertEqual(self.store.config['model'],'old-saved-id')
        self.assertEqual((self.store.folder/'settings.json').read_bytes(),before)
        self.assertEqual((self.store.folder/'credential.bin').read_bytes(),credential)
        with self.assertRaises(ValueError):self.store.model_settings({'base_url':'https://example.invalid'})
        config,key=self.store.model_settings({'base_url':'https://example.invalid','api_key':'SYNTHETIC-OTHER','model':'DeepSeek V4.1 Flash'})
        self.assertEqual(config['model'],'DeepSeek V4.1 Flash');self.assertEqual(key,'SYNTHETIC-OTHER')
    def test_connection_endpoint_tests_draft_and_never_returns_key(self):
        self.store.save({'api_key':'SYNTHETIC-LOCAL-ONLY','model':'saved-old'})
        server=Server(self.store);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        def post(data):
            req=urllib.request.Request(server.base+'check-model',data=json.dumps(data).encode(),headers={'Content-Type':'application/json'})
            return json.load(urllib.request.urlopen(req))
        try:
            with patch('server.completion',return_value={'content':'连接成功'}) as call:
                result=post({'model':'DeepSeek V4.1 Flash','api_key':'SYNTHETIC-DRAFT'})
                self.assertEqual(result['model'],'deepseek-flash');self.assertTrue(result['ok'])
                self.assertEqual(call.call_args.args[0]['model'],'deepseek-flash');self.assertEqual(call.call_args.args[1],'SYNTHETIC-DRAFT')
                self.assertEqual(call.call_args.kwargs['max_tokens'],64)
                self.assertNotIn('SYNTHETIC',json.dumps(result));self.assertEqual(self.store.key,'SYNTHETIC-LOCAL-ONLY')
                self.assertEqual(self.store.config['model'],'saved-old');self.assertFalse(self.store.config['use_model'])
            with patch('server.completion',side_effect=ValueError('API余额不足（HTTP 402）')):
                with self.assertRaises(urllib.error.HTTPError) as failure:post({'model':'deepseek-flash'})
                self.assertIn('402',json.load(failure.exception)['error']);self.assertEqual(self.store.config['model'],'saved-old')
        finally:server.shutdown();server.server_close();thread.join(3)
    def test_supported_readers_and_sources(self):
        from docx import Document
        from openpyxl import Workbook
        from pypdf import PdfWriter
        from pypdf.generic import NameObject,DictionaryObject,DecodedStreamObject
        root=Path(self.store.config['workspace'])
        doc=Document();doc.add_paragraph('Synthetic deadline 2099-01-08');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Test';table.cell(0,1).text='100';doc.save(root/'synthetic.docx')
        book=Workbook();book.active['A1']='Synthetic';book.active['C2']='=1+1';book.save(root/'synthetic.xlsx');book.close()
        writer=PdfWriter();page=writer.add_blank_page(width=300,height=400)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        fonts=DictionaryObject({NameObject('/F1'):font})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):fonts})
        stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 30 300 Td (Synthetic deadline 2099-01-08) Tj ET');page[NameObject('/Contents')]=writer._add_object(stream)
        with (root/'synthetic.pdf').open('wb') as file:writer.write(file)
        self.assertIn('2099',self.tools.read_document('synthetic.docx')['text'])
        rows=self.tools.read_document('synthetic.xlsx')['rows'];self.assertTrue(rows[1][2]['formula']);self.assertEqual(rows[1][0]['value'],'')
        self.assertIn('2099',self.tools.read_document('synthetic.pdf')['text'])
        self.assertEqual(self.tools.read_document('示例会议记录.md')['source']['sha256'],hashlib.sha256((root/'示例会议记录.md').read_bytes()).hexdigest())
    def test_scope_and_sensitive_files(self):
        with self.assertRaises(ValueError):self.tools.read_document('../settings.json')
        with self.assertRaises(ValueError):self.tools.read_document(str(self.store.folder/'settings.json'))
        root=Path(self.store.config['workspace']);(root/'.env').write_text('SYNTHETIC SECRET')
        self.assertNotIn('.env',[f['path'] for f in self.tools.list_files()['files']])
        with self.assertRaises(ValueError):self.store.choose_workspace('')
        with self.assertRaises(ValueError):self.store.choose_workspace(str(Path(root.anchor)))
    def test_search_and_exact_decimal(self):
        self.assertTrue(self.tools.search_documents('2099-01-08')['matches'])
        self.assertEqual(calculate_decimal('sum([100.50,20.25,30.00])'),'150.75')
        for expression in ["__import__('os').system('anything')",'1/0','2**1000000']:
            with self.assertRaises(ValueError):calculate_decimal(expression)
    def test_real_outputs_no_overwrite_and_no_formula_execution(self):
        from docx import Document
        from openpyxl import load_workbook
        source=Path(self.store.config['workspace'])/'示例会议记录.md';before=source.read_bytes()
        a=self.tools.create_document('纪要','合成测试','# 待确认\n会议时间待确认')
        b=self.tools.create_document('纪要','另一个标题','合成正文')
        self.assertNotEqual(a['name'],b['name']);self.assertTrue(any('待确认' in p.text for p in Document(self.tools.output/a['name']).paragraphs))
        x=self.tools.create_workbook('费用',[{'title':'合成','columns':['值','说明'],'rows':[[150.75,'=HYPERLINK("https://example.com")']]}])
        book=load_workbook(self.tools.output/x['name']);self.assertEqual(book.active['A2'].value,150.75);self.assertNotEqual(book.active['B2'].data_type,'f');book.close()
        self.assertEqual(source.read_bytes(),before)
        with self.assertRaises(ValueError):self.tools.create_document('../outside','title','text')
    def test_memory_is_explicit_preferences_and_deletable(self):
        self.tools.remember_preference('报告风格','先给结论，再列依据')
        self.assertIn('先给结论',Store(self.store.folder).preferences['报告风格'])
        with self.assertRaises(ValueError):self.tools.remember_preference('客户欠款','100')
        with self.assertRaises(ValueError):self.tools.remember_preference('工作习惯','我的API密钥是SYNTHETIC')
        self.store.forget('报告风格');self.assertFalse(Store(self.store.folder).preferences)
    def test_existing_artifact_can_be_read_and_revised_without_overwrite(self):
        first=self.tools.create_document('合成纪要','合成纪要','原版内容与待确认问题')
        original=(self.tools.output/first['name']).read_bytes()
        self.assertEqual(self.tools.list_artifacts()['artifacts'][0]['id'],first['id'])
        content=self.tools.read_artifact(first['id']);self.assertIn('原版内容',content['text']);self.assertEqual(content['source']['artifact_id'],first['id'])
        second=self.tools.create_document('合成纪要修订','合成纪要修订','改进后的内容，保留未知问题')
        self.assertNotEqual(first['id'],second['id']);self.assertEqual((self.tools.output/first['name']).read_bytes(),original)
        self.assertEqual(Store(self.store.folder).config['workspace_id'],self.store.config['workspace_id'])
    def test_local_demo_does_not_apply_stale_template_to_changed_files(self):
        root=Path(self.store.config['workspace']);(root/'示例费用.csv').write_text('完全不同的数据',encoding='utf-8')
        result=OfficeAgent().run('演示整理示例会议资料','office',self.store.config,'',[],lambda *a:None,self.store)
        self.assertFalse(result['artifacts']);self.assertIn('已停止套用',result['body'])
    def test_local_demo_creates_real_artifacts_from_unchanged_demo(self):
        result=OfficeAgent().run('演示整理示例会议资料','office',self.store.config,'',[],lambda *a:None,self.store)
        self.assertEqual(len(result['artifacts']),2);self.assertIn('未调用模型',result['engine'])
    def test_real_framework_multi_turn_loop(self):
        model=ThreadingHTTPServer(('127.0.0.1',0),ScriptedModel);model.count=0;model.payloads=[];model.delay=0
        model.script=[('list_files',{}),('read_document',{'path':'示例会议记录.md'}),('calculate',{'expression':'100.5+20.25+30'}),
            ('create_document',{'name':'合成纪要','title':'合成纪要','content':'# 待办\n示例负责人甲整理清单。'}),
            ('create_workbook',{'name':'合成费用','sheets':[{'title':'费用','columns':['项目','金额'],'rows':[['合成合计',150.75]]}]})]
        thread=threading.Thread(target=model.serve_forever,daemon=True);thread.start()
        try:
            config=dict(self.store.config,base_url=f'http://127.0.0.1:{model.server_port}',model='synthetic-office-stub',use_model=True)
            result=OfficeAgent().run('处理合成资料并生成两个成果','office',config,'SYNTHETIC-TEST',[],lambda *a:None,self.store)
            self.assertEqual(result['kind'],'office');self.assertEqual(model.count,6);self.assertEqual(len(result['artifacts']),2)
            self.assertTrue(any(m['role']=='tool' for m in model.payloads[-1]['messages']))
        finally:model.shutdown();model.server_close();thread.join(3)
    def test_http_artifact_and_scope_change_protection(self):
        server=Server(self.store);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        def api(path,data=None):
            req=server.base+path if data is None else urllib.request.Request(server.base+path,data=json.dumps(data).encode(),headers={'Content-Type':'application/json'})
            return urllib.request.urlopen(req)
        try:
            artifact=self.tools.create_document('合成成果','合成','test')
            self.assertEqual(hashlib.sha256(api('artifact/'+artifact['id']).read()).hexdigest(),artifact['sha256'])
            with self.assertRaises(urllib.error.HTTPError):api('artifact/../../settings.json')
            with self.assertRaises(urllib.error.HTTPError):urllib.request.urlopen(server.origin+'/status')
            req=urllib.request.Request(server.base+'ask',data=b'{}',headers={'Origin':'https://example.com'})
            with self.assertRaises(urllib.error.HTTPError):urllib.request.urlopen(req)
            self.assertNotIn('api_key',api('status').read().decode())
        finally:server.shutdown();server.server_close();thread.join(3)
    def test_cancel_interrupts_model_and_workspace_cannot_change_mid_task(self):
        model=ThreadingHTTPServer(('127.0.0.1',0),ScriptedModel);model.count=0;model.payloads=[];model.delay=3;model.script=[]
        mt=threading.Thread(target=model.serve_forever,daemon=True);mt.start()
        self.store.save({'base_url':f'http://127.0.0.1:{model.server_port}','model':'synthetic-office-stub','api_key':'SYNTHETIC-CANCEL-TEST','use_model':True})
        server=Server(self.store);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        def post(path,data):
            req=urllib.request.Request(server.base+path,data=json.dumps(data).encode(),headers={'Content-Type':'application/json'})
            return json.load(urllib.request.urlopen(req))
        try:
            jid=post('ask',{'question':'等待模型回复的合成测试','mode':'office'})['id'];time.sleep(.25)
            with self.assertRaises(urllib.error.HTTPError):post('workspace',{'path':self.store.config['workspace']})
            post('cancel',{'id':jid});started=time.monotonic()
            while time.monotonic()-started<2:
                job=json.load(urllib.request.urlopen(server.base+'job/'+jid))
                if job['state']=='cancelled':break
                time.sleep(.05)
            self.assertEqual(job['state'],'cancelled');self.assertFalse(job['answer']['artifacts'])
        finally:server.shutdown();server.server_close();thread.join(3);model.shutdown();model.server_close();mt.join(3)

if __name__=='__main__':unittest.main(verbosity=2)
